"""导出服务测试：TXT/CSV/XLSX，覆盖空数据、多账号、中文、特殊字符、编码、去重、过滤。"""
from __future__ import annotations

import csv
import os
import tempfile
import unittest

from app.export.export_service import ExportService
from app.export.models import AccountStat, ExportRow


def _row(video_id, username="user_a", account="acc_a", publish_date="2026-09-24",
         task_id="TASK-1", video_url=None, **kw):
    return ExportRow(
        account=account, username=username, video_id=video_id,
        publish_time="2026-09-24 18:30:00", publish_date=publish_date,
        video_url=video_url or f"https://www.tiktok.com/@{username}/video/{video_id}",
        collect_time="2026-09-24 20:00:00", task_id=task_id, status="ok", **kw)


class TestExport(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = self.tmp.name

    def tearDown(self):
        self.tmp.cleanup()

    def _path(self, name):
        return os.path.join(self.dir, name)

    # 1. 空数据导出
    def test_empty_export(self):
        ExportService.export_txt([], self._path("e.txt"))
        ExportService.export_csv([], self._path("e.csv"))
        ExportService.export_xlsx([], [], self._path("e.xlsx"))
        self.assertTrue(os.path.exists(self._path("e.txt")))
        self.assertTrue(os.path.exists(self._path("e.csv")))
        self.assertTrue(os.path.exists(self._path("e.xlsx")))

    # 2. 单账号导出
    def test_single_account(self):
        rows = [_row("1"), _row("2")]
        ExportService.export_csv(rows, self._path("s.csv"))
        with open(self._path("s.csv"), encoding="utf-8-sig") as f:
            lines = f.read().strip().splitlines()
        self.assertEqual(len(lines), 3)  # header + 2

    # 3. 多账号导出
    def test_multi_account(self):
        rows = [_row("1", username="a"), _row("2", username="b"), _row("3", username="c")]
        ExportService.export_csv(rows, self._path("m.csv"))
        with open(self._path("m.csv"), encoding="utf-8-sig") as f:
            lines = f.read().strip().splitlines()
        self.assertEqual(len(lines), 4)

    # 4. 中文账号/备注
    def test_chinese(self):
        rows = [_row("1", username="中文用户", account="账号甲")]
        ExportService.export_csv(rows, self._path("c.csv"))
        with open(self._path("c.csv"), encoding="utf-8-sig") as f:
            text = f.read()
        self.assertIn("中文用户", text)
        self.assertIn("账号甲", text)

    # 5. URL 特殊字符
    def test_url_special_chars(self):
        url = "https://www.tiktok.com/@u/video/1?is_from_webapp=1&sender_device=pc"
        rows = [_row("1", video_url=url)]
        ExportService.export_csv(rows, self._path("u.csv"))
        with open(self._path("u.csv"), encoding="utf-8-sig") as f:
            text = f.read()
        self.assertIn(url, text)

    # 5b. CSV 特殊字符（逗号/引号/换行）转义
    def test_csv_special_characters(self):
        rows = [_row("1", username='用户"引号",逗号\n换行')]
        ExportService.export_csv(rows, self._path("sc.csv"))
        with open(self._path("sc.csv"), encoding="utf-8-sig", newline="") as f:
            data = list(csv.reader(f))
        self.assertEqual(data[1][1], '用户"引号",逗号\n换行')  # 字段还原

    # 6. CSV 编码（BOM）
    def test_csv_encoding_bom(self):
        ExportService.export_csv([_row("1")], self._path("bom.csv"))
        with open(self._path("bom.csv"), "rb") as f:
            self.assertEqual(f.read(3), b"\xef\xbb\xbf")

    # 7. XLSX 文件可正常打开
    def test_xlsx_openable(self):
        rows = [_row("1"), _row("2")]
        stats = [AccountStat(account="acc_a", username="user_a",
                             target_count=4, actual_count=2, status="partial")]
        ExportService.export_xlsx(rows, stats, self._path("x.xlsx"))
        from openpyxl import load_workbook
        wb = load_workbook(self._path("x.xlsx"))
        self.assertEqual(wb.sheetnames, ["作品链接", "账号统计"])
        ws1 = wb["作品链接"]
        self.assertEqual(ws1.max_row, 3)  # header + 2
        ws2 = wb["账号统计"]
        self.assertEqual(ws2.max_row, 2)  # header + 1

    # 8. TXT 纯 URL
    def test_txt_urls(self):
        rows = [_row("1", username="a"), _row("2", username="b")]
        ExportService.export_txt(rows, self._path("u.txt"), mode="urls")
        with open(self._path("u.txt"), encoding="utf-8") as f:
            lines = [x for x in f.read().split("\n") if x.strip()]
        self.assertEqual(lines, ["https://www.tiktok.com/@a/video/1",
                                 "https://www.tiktok.com/@b/video/2"])

    # 8b. TXT 分组
    def test_txt_grouped(self):
        rows = [_row("1", username="a"), _row("2", username="a"), _row("3", username="b")]
        ExportService.export_txt(rows, self._path("g.txt"), mode="grouped")
        with open(self._path("g.txt"), encoding="utf-8") as f:
            text = f.read()
        self.assertIn("@a\nhttps://www.tiktok.com/@a/video/1\n\nhttps://www.tiktok.com/@a/video/2", text)
        self.assertIn("@b\nhttps://www.tiktok.com/@b/video/3", text)

    # 9. 去重后导出
    def test_dedupe_export(self):
        rows = [_row("1"), _row("1"), _row("2")]  # video_id 1 重复
        deduped = ExportService.dedupe_rows(rows)
        self.assertEqual(len(deduped), 2)
        ExportService.export_csv(deduped, self._path("d.csv"))
        with open(self._path("d.csv"), encoding="utf-8-sig") as f:
            lines = f.read().strip().splitlines()
        self.assertEqual(len(lines), 3)

    # 10. 按任务导出
    def test_filter_by_task(self):
        rows = [_row("1", task_id="T1"), _row("2", task_id="T2"), _row("3", task_id="T1")]
        filtered = ExportService.filter_rows(rows, task_id="T1")
        self.assertEqual([r.video_id for r in filtered], ["1", "3"])

    # 11. 按日期导出
    def test_filter_by_date(self):
        rows = [_row("1", publish_date="2026-09-24"), _row("2", publish_date="2026-09-23")]
        filtered = ExportService.filter_rows(rows, date="2026-09-24")
        self.assertEqual([r.video_id for r in filtered], ["1"])


if __name__ == "__main__":
    unittest.main()
